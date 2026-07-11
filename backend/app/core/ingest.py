"""Document text extraction for onboarding uploads (C8 / Sprint 7).

Turns an uploaded PDF / Word / PowerPoint / Excel / CSV / Markdown / text file
into plain text that gets written into the memory system as onboarding
context. Deliberately simple: text only, no images/layout/formulas, with a
size cap so a giant file can't blow up a memory record. Legacy pre-2007
Office formats (.ppt/.xls) are out of scope — users can re-save as the
modern format.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass

MAX_CHARS = 100_000
# Belt-and-braces under MAX_CHARS: a huge spreadsheet stops contributing rows
# long before the char cap silently eats every later sheet.
MAX_ROWS_PER_SHEET = 1_000
SUPPORTED_EXTENSIONS = (
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".csv",
    ".tsv",
    ".md",
    ".markdown",
    ".txt",
)


class UnsupportedFileType(ValueError):
    pass


@dataclass
class ParsedDocument:
    title: str
    text: str


def _strip_ext(filename: str) -> str:
    name = filename.rsplit("/", 1)[-1]
    for ext in SUPPORTED_EXTENSIONS:
        if name.lower().endswith(ext):
            return name[: -len(ext)]
    return name


def _parse_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p.strip() for p in parts if p.strip())


def _parse_docx(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def _shape_texts(shape) -> list[str]:
    parts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    for sub in getattr(shape, "shapes", ()):  # grouped shapes nest
        parts.extend(_shape_texts(sub))
    return parts


def _parse_pptx(content: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    slides: list[str] = []
    for idx, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_texts(shape))
        if slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            notes = notes_frame.text.strip() if notes_frame is not None else ""
            if notes:
                parts.append(f"Notes: {notes}")
        if parts:
            slides.append(f"Slide {idx}:\n" + "\n".join(parts))
    return "\n\n".join(slides)


def _parse_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    # data_only gives the last-computed value for formula cells (None when the
    # file was never opened in Excel — acceptable for text extraction).
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheets: list[str] = []
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if not cells or not any(cells):
                    continue
                rows.append(" | ".join(cells))
                if len(rows) >= MAX_ROWS_PER_SHEET:
                    rows.append(f"[truncated at {MAX_ROWS_PER_SHEET} rows]")
                    break
            if rows:
                sheets.append(f"Sheet: {worksheet.title}\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        workbook.close()


def _parse_delimited(content: bytes, delimiter: str) -> str:
    text = content.decode("utf-8", errors="replace")
    rows: list[str] = []
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        cells = [c.strip() for c in row]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def parse_upload(filename: str, content: bytes) -> ParsedDocument:
    """Extract text from an uploaded file. Raises UnsupportedFileType for
    anything outside SUPPORTED_EXTENSIONS."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = _parse_pdf(content)
    elif lower.endswith(".docx"):
        text = _parse_docx(content)
    elif lower.endswith(".pptx"):
        text = _parse_pptx(content)
    elif lower.endswith(".xlsx"):
        text = _parse_xlsx(content)
    elif lower.endswith(".csv"):
        text = _parse_delimited(content, delimiter=",")
    elif lower.endswith(".tsv"):
        text = _parse_delimited(content, delimiter="\t")
    elif lower.endswith((".md", ".markdown", ".txt")):
        text = content.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFileType(
            f"Unsupported file type for '{filename}'. "
            f"Supported: {', '.join(SUPPORTED_EXTENSIONS)}."
        )

    text = text.strip()[:MAX_CHARS]
    return ParsedDocument(title=_strip_ext(filename), text=text)
